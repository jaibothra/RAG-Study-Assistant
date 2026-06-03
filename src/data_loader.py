from pathlib import Path
from typing import Any, List

from langchain_community.document_loaders import (
    CSVLoader,
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".csv", ".xlsx", ".docx", ".json", ".md"}


def iter_data_files(data_dir: Path):
    """Yield supported document paths under data_dir."""
    data_path = data_dir.resolve()
    if not data_path.exists():
        return
    for file_path in data_path.rglob("*"):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield file_path


def latest_data_mtime(data_dir: Path) -> float:
    """Latest modification time of any supported file in data_dir, or 0 if none."""
    latest = 0.0
    for file_path in iter_data_files(data_dir):
        latest = max(latest, file_path.stat().st_mtime)
    return latest


def _loaded_label(ext: str, count: int) -> str:
    if ext == ".pdf":
        unit = "page" if count == 1 else "pages"
    elif ext == ".csv":
        unit = "row" if count == 1 else "rows"
    else:
        unit = "part" if count == 1 else "parts"
    return f"{count} {unit}"


def _load_xlsx(file_path: Path) -> List[Document]:
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts: List[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    text = "\n".join(parts)
    return [Document(page_content=text, metadata={"source": str(file_path)})]


def load_all_documents(data_dir: str) -> List[Any]:
    """
    Load supported files from data_dir into LangChain Document objects.
    Supported: PDF, TXT, MD, CSV, XLSX, DOCX, JSON (as plain text)
    """
    data_path = Path(data_dir).resolve()
    if not data_path.exists():
        raise FileNotFoundError(f"Data directory not found: {data_path}")

    documents: List[Any] = []
    file_count = 0
    loaders = {
        ".pdf": lambda p: PyPDFLoader(str(p)),
        ".txt": lambda p: TextLoader(str(p), encoding="utf-8"),
        ".md": lambda p: TextLoader(str(p), encoding="utf-8"),
        ".csv": lambda p: CSVLoader(str(p)),
        ".docx": lambda p: Docx2txtLoader(str(p)),
        ".json": lambda p: TextLoader(str(p), encoding="utf-8"),
    }

    for file_path in sorted(iter_data_files(data_path)):
        ext = file_path.suffix.lower()
        try:
            if ext == ".xlsx":
                loaded = _load_xlsx(file_path)
            else:
                loaded = loaders[ext](file_path).load()
            documents.extend(loaded)
            file_count += 1
            print(f"[OK] {file_path.name}: {_loaded_label(ext, len(loaded))}")
        except Exception as e:
            print(f"[WARN] Skipped {file_path}: {e}")

    print(
        f"[INFO] Ingested {file_count} file(s), {len(documents)} pages/rows/parts "
        f"(chunking happens next)"
    )
    return documents
